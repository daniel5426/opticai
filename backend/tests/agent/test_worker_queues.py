"""Requested exports cannot be blocked by summaries or consume another clinic."""
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from agent.config import now, data_directory
from agent.daily import enqueue_daily
from agent.models import AgentJob
from agent.retention import cleanup_expired
from models import User
from workers import agent_worker as worker


def add_job(db, auth, key, kind, **extra):
    job = AgentJob(unique_key=key, company_id=auth.user.company_id,
        clinic_id=auth.clinic_id, user_id=auth.user.id, context_id=auth.context_id,
        kind=kind, status='queued', **extra)
    db.add(job)
    db.commit()
    return job


def test_report_claim_skips_older_brief_and_foreign_report(clinic):
    db, auth, patients, _ = clinic
    brief = add_job(db, auth, 'brief', 'brief', created_at=now() - timedelta(days=1))
    foreign = add_job(db, auth, 'foreign', 'report', created_at=now() - timedelta(hours=1))
    foreign.clinic_id = patients[2].clinic_id
    db.commit()
    report = add_job(db, auth, 'report', 'report')
    assert worker.claim(db, 'reports', 'requested', [auth.clinic_id]).id == report.id
    db.refresh(brief)
    db.refresh(foreign)
    assert brief.status == foreign.status == 'queued'


@pytest.mark.parametrize('kind', ['report', 'document'])
def test_summary_claim_never_consumes_requested_work(clinic, kind):
    db, auth, *_ = clinic
    requested = add_job(db, auth, kind, kind)
    brief = add_job(db, auth, 'brief', 'brief')
    assert worker.claim(db, 'summaries', 'summaries', [auth.clinic_id]).id == brief.id
    db.refresh(requested)
    assert requested.status == 'queued'


def test_expired_leases_keep_queue_and_clinic_boundaries(clinic):
    db, auth, patients, _ = clinic
    rows = [add_job(db, auth, key, kind) for key, kind in
            [('expired-brief', 'brief'), ('foreign', 'report'), ('expired-report', 'report')]]
    for row in rows:
        row.status, row.lease_until = 'running', now() - timedelta(minutes=1)
    rows[1].clinic_id = patients[2].clinic_id
    db.commit()
    assert worker.claim(db, 'new-reports', 'requested', [auth.clinic_id]).id == rows[2].id
    assert worker.claim(db, 'new-reports', 'requested', [auth.clinic_id]) is None


def test_daily_admission_is_limited_and_idempotent(clinic):
    db, auth, patients, _ = clinic
    db.add(User(company_id=patients[2].company_id, clinic_id=patients[2].clinic_id,
                username='foreign-staff', role_level=3, is_active=True))
    db.commit()
    at = datetime(2026, 10, 25, 8, tzinfo=timezone.utc)
    assert enqueue_daily(db, at, clinic_ids=[auth.clinic_id]) == 3
    assert enqueue_daily(db, at, clinic_ids=[auth.clinic_id]) == 0
    assert {row.clinic_id for row in db.query(AgentJob)} == {auth.clinic_id}
    assert enqueue_daily(db, at, clinic_ids=[]) == 0


def test_requested_worker_exports_without_scheduling_summaries(clinic, monkeypatch):
    db, auth, _, sessions = clinic
    monkeypatch.setattr(worker, 'SessionLocal', sessions)
    monkeypatch.setattr(worker.time, 'time', lambda: 1_800_000_010)
    def unexpected_scheduler(*args, **kwargs):
        pytest.fail('Requested worker must not enqueue daily summaries')
    monkeypatch.setattr(worker, 'enqueue_daily', unexpected_scheduler)
    add_job(db, auth, 'older-brief', 'brief', created_at=now() - timedelta(days=1))
    job = add_job(db, auth, 'export', 'report', definition={
        'dataset': 'clients', 'columns': ['id'], 'format': 'xlsx'})
    assert worker.run_once('requested-test', 'requested', [auth.clinic_id])
    db.refresh(job)
    assert job.status == 'completed' and job.result['rows'] == 2
    workbook = load_workbook(data_directory() / job.file_name)
    assert workbook['Records'].max_row == 3
    workbook.close()


def test_scoped_retention_does_not_change_foreign_jobs(clinic):
    db, auth, patients, _ = clinic
    local = add_job(db, auth, 'local-upload', 'document', created_at=now() - timedelta(hours=2))
    foreign = add_job(db, auth, 'foreign-upload', 'document', created_at=now() - timedelta(hours=2))
    local.status = foreign.status = 'uploading'
    foreign.clinic_id = patients[2].clinic_id
    db.commit()
    cleanup_expired(db, clinic_ids=[auth.clinic_id])
    db.refresh(local)
    db.refresh(foreign)
    assert local.status == 'failed' and foreign.status == 'uploading'
