from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from models import Order, Client, AuthSession, Appointment, OpticalExam, ExamLayout, ExamLayoutInstance
from agent import reports, records, changes
from agent.authority import Authority, resolve_artifact
from agent.config import now, data_directory
from agent.models import AgentJob, FollowupTask, TaskEvent
from agent.daily import enqueue_daily
from services.appointment_schedule import available, find_slots


def definition(**extra):
    return {'dataset': 'clients', 'as_of': '2026-10-14', 'filters': [
        {'field': 'gender', 'operator': 'eq', 'value': 'male'},
        {'field': 'age', 'operator': 'lt', 'value': 18},
        {'field': 'outstanding', 'operator': 'gt', 'value': 0}], **extra}


def test_age_gender_balance_and_shared_invoice(clinic):
    db, auth, patients, _ = clinic
    result = reports.preview(db, auth, definition())
    assert result['total'] == 1
    row = result['rows'][0]
    assert row['age'] == 12
    assert row['total'] == 500 and row['paid'] == 100 and row['outstanding'] == 400
    assert row['id'] == patients[0].id


def test_birthday_boundary_unknown_age_and_scope(clinic):
    db, auth, patients, _ = clinic
    patients[0].date_of_birth = date(2008, 10, 14); db.commit()
    assert reports.preview(db, auth, definition())['total'] == 0
    patients[0].date_of_birth = date(2008, 10, 15); db.commit()
    assert reports.preview(db, auth, definition())['total'] == 1
    patients[0].date_of_birth = None; db.commit()
    assert reports.preview(db, auth, definition())['total'] == 0
    result = reports.preview(db, auth, {'dataset': 'clients', 'columns': ['id']})
    assert patients[2].id not in [row['id'] for row in result['rows']]


def test_report_fields_and_sql_rejected(clinic):
    db, auth, *_ = clinic
    for field in ['hidden_note', 'google_access_token', 'drop table clients', 'password_hash']:
        with pytest.raises(HTTPException):
            reports.preview(db, auth, {'dataset': 'clients', 'columns': [field]})


def test_full_export_and_formula_safety(clinic):
    db, auth, patients, _ = clinic
    for index in range(35):
        db.add(Client(company_id=auth.user.company_id, clinic_id=auth.clinic_id, first_name='=1+1', last_name=str(index)))
    db.commit()
    raw = {'dataset': 'clients', 'columns': ['first_name', 'last_name']}
    assert len(reports.preview(db, auth, raw)['rows']) == 25
    queued = reports.queue_export(db, auth, raw, 'stable-export')
    assert reports.queue_export(db, auth, raw, 'stable-export')['job_id'] == queued['job_id']
    job = db.get(AgentJob, queued['job_id'])
    reports.export_report(db, auth, job)
    workbook = load_workbook(data_directory() / job.file_name)
    assert workbook['Records'].max_row == 38
    assert workbook['Records']['A4'].data_type == 's'
    assert reports.safe_csv('=1+1') == "'=1+1"


def test_immutable_change_stale_and_idempotent_receipt(clinic):
    db, auth, patients, _ = clinic
    result = changes.prepare(db, auth, 'clients', 'update', {'notes': 'Reviewed'}, patients[0].id)
    changes.approve(db, auth, result['change_id'])
    assert patients[0].notes is None
    changed = changes.commit(db, auth, result['change_id'])
    assert changed['notes'] == 'Reviewed'
    assert changes.commit(db, auth, result['change_id']) == changed
    result = changes.prepare(db, auth, 'clients', 'update', {'notes': 'New'}, patients[0].id)
    changes.approve(db, auth, result['change_id'])
    patients[0].notes = 'Concurrent edit'; db.commit()
    with pytest.raises(HTTPException) as error:
        changes.commit(db, auth, result['change_id'])
    assert error.value.status_code == 409


def test_foreign_patient_and_revoked_session(clinic):
    db, auth, patients, _ = clinic
    with pytest.raises(HTTPException):
        records.history(db, auth, patients[2].id)
    context = resolve_artifact(db, {'external_session_id': auth.context_id})
    assert context.user.id == auth.user.id
    session = db.query(AuthSession).first(); session.revoked_at = now(); db.commit()
    with pytest.raises(HTTPException):
        resolve_artifact(db, {'external_session_id': auth.context_id})


def test_patient_context_cannot_retarget(clinic):
    db, auth, patients, _ = clinic
    auth.patient_id = patients[0].id
    with pytest.raises(HTTPException):
        records.history(db, auth, patients[1].id)


def test_tasks_assign_complete_audit(clinic):
    db, auth, patients, _ = clinic
    proposed = changes.prepare(db, auth, 'tasks', 'create', {'client_id': patients[0].id, 'title': 'Follow up', 'due_date': '2026-10-15'})
    changes.approve(db, auth, proposed['change_id'])
    result = changes.commit(db, auth, proposed['change_id'])
    assert db.query(FollowupTask).count() == 1 and db.query(TaskEvent).count() == 1
    proposed = changes.prepare(db, auth, 'tasks', 'update', {'status': 'completed'}, result['id'])
    changes.approve(db, auth, proposed['change_id'])
    result = changes.commit(db, auth, proposed['change_id'])
    assert result['status'] == 'completed' and result['revision'] == 2


def test_duration_conflicts_breaks_and_slots(clinic):
    db, auth, patients, _ = clinic
    day = date(2026, 11, 8)
    db.add(Appointment(client_id=patients[0].id, clinic_id=auth.clinic_id, user_id=auth.user.id, date=day, time='09:00', duration=60)); db.commit()
    assert not available(db, auth.clinic_id, auth.user.id, day, '09:45', 30)
    assert not available(db, auth.clinic_id, auth.user.id, day, '11:45', 30)
    assert available(db, auth.clinic_id, auth.user.id, day, '10:00', 30)
    result = find_slots(db, auth, day.isoformat(), day.isoformat())
    assert all(row['time'] != '09:30' for row in result['slots'])


def test_exam_measurements_and_optical_equivalence(clinic):
    db, auth, patients, _ = clinic
    layout = ExamLayout(clinic_id=auth.clinic_id, name='Exam', layout_data='{}'); db.add(layout); db.flush()
    exam = OpticalExam(client_id=patients[0].id, clinic_id=auth.clinic_id, exam_date=date.today()); db.add(exam); db.flush()
    db.add(ExamLayoutInstance(exam_id=exam.id, layout_id=layout.id, exam_data={'final-prescription': {'r_sph': -2}})); db.commit()
    assert records.history(db, auth, patients[0].id)['exam_layouts'][0]['data']['final-prescription']['r_sph'] == -2
    assert records.prescription({'r_sph': -1, 'r_cyl': 1, 'r_ax': 0}) == records.prescription({'r_sph': 0, 'r_cyl': -1, 'r_ax': 90})


def test_daily_six_am_unique_and_dst(clinic):
    db, *_ = clinic
    zone = ZoneInfo('Asia/Jerusalem')
    assert enqueue_daily(db, datetime(2026, 10, 25, 5, 59, tzinfo=zone)) == 0
    assert enqueue_daily(db, datetime(2026, 10, 25, 6, 0, tzinfo=zone)) == 3
    assert enqueue_daily(db, datetime(2026, 10, 25, 6, 1, tzinfo=zone)) == 0


def test_aggregates_without_groups_are_totals(clinic):
    db, auth, clients, _ = clinic
    from agent.reports import preview
    result = preview(db, auth, {'dataset': 'clients', 'aggregates': [{'field': 'outstanding', 'operation': 'sum'}]})
    ils = next(row for row in result['rows'] if row['currency'] == 'ILS')
    assert ils['sum_outstanding'] == 400
    assert 'id' not in result['columns']


def test_deleted_records_are_not_agent_evidence(clinic):
    db, auth, clients, _ = clinic
    from agent.records import list_records
    from agent.config import now
    row = db.query(Order).first(); row.deleted_at = now(); db.commit()
    assert list_records(db, auth, 'orders')['total'] == 0


def test_order_comparison_preserves_eye_and_transposition(clinic):
    from agent.records import prescription_comparison
    result = prescription_comparison({'r_sph': -1, 'r_cyl': 1, 'r_ax': 90}, {'r_sph': 0, 'r_cyl': -1, 'r_ax': 0})
    assert result['eyes']['r']['equivalent'] is True
    assert result['eyes']['l']['comparable'] is False


def test_commit_rechecks_role_after_proposal(clinic):
    db, auth, clients, _ = clinic
    from agent.changes import prepare, commit
    from fastapi import HTTPException
    proposal = prepare(db, auth, 'orders', 'update', {'client_id': clients[0].id, 'type': 'glasses'}, record_id=db.query(Order).first().id)
    changes.approve(db, auth, proposal['change_id'])
    auth.user.role_level = 1; db.commit()
    with pytest.raises(HTTPException) as error:
        commit(db, auth, proposal['change_id'])
    assert error.value.status_code == 403


def test_task_assignment_stays_in_selected_clinic(clinic):
    db, auth, clients, _ = clinic
    from models import User, Clinic
    from agent.changes import prepare
    from fastapi import HTTPException
    other = Clinic(company_id=auth.user.company_id, name='Other clinic', unique_id='another');db.add(other);db.flush()
    user = User(company_id=auth.user.company_id, clinic_id=other.id, username='otherstaff', role_level=3);db.add(user);db.commit()
    with pytest.raises(HTTPException):
        prepare(db, auth, 'tasks', 'create', {'assigned_user_id': user.id, 'title': 'Follow up', 'due_date': '2026-10-15'})


def test_document_job_is_not_claimable_until_upload_completes(clinic, monkeypatch):
    db, auth, clients, _ = clinic
    from agent.documents import submit_document
    from agent.models import AgentJob
    class Uploads:
        def create_upload(self, *args, **kwargs):
            assert db.query(AgentJob).one().status == 'uploading'
            return {'id': 'upload', 'part_bytes': 1024, 'parts': [{'part_number': 1, 'completed': False, 'upload_url': '/local'}]}
        def put_upload_part(self, *args): pass
        def complete_upload(self, *args, **kwargs): return {'input_ref': 'pipeline-input://' + 'a' * 32}
    from types import SimpleNamespace
    monkeypatch.setattr('agent.documents.client', lambda: SimpleNamespace(ingestion_jobs=Uploads()))
    result = submit_document(db, auth, 'visit.txt', 'text/plain', b'Synthetic visit')
    assert result['status'] == 'queued'
    assert db.query(AgentJob).one().definition['input_ref'].startswith('pipeline-input://')


def test_refresh_preserves_context_but_logout_revokes_it(clinic):
    db, auth, clients, _ = clinic
    from agent.models import AgentContext
    from agent.authority import rotate_session_contexts
    previous = db.get(AuthSession, db.get(AgentContext, auth.context_id).session_id)
    successor = AuthSession(id='00000000-0000-0000-0000-000000000002', user_id=auth.user.id,
        company_id=auth.user.company_id, clinic_id=auth.clinic_id, refresh_token_hash='successor', expires_at=now() + timedelta(days=1))
    db.add(successor);db.flush();previous.revoked_at = now()
    rotate_session_contexts(db, previous, successor.id);db.commit()
    assert resolve_artifact(db, {'external_session_id': auth.context_id}).user.id == auth.user.id
    successor.revoked_at = now();db.commit()
    with pytest.raises(HTTPException):
        resolve_artifact(db, {'external_session_id': auth.context_id})


def test_commit_is_fail_closed_without_product_approval(clinic):
    db, auth, patients, _ = clinic
    proposal = changes.prepare(db, auth, 'clients', 'update', {'notes': 'Must not be saved'}, patients[0].id)
    with pytest.raises(HTTPException) as error:
        changes.commit(db, auth, proposal['change_id'])
    assert error.value.detail == 'agent.approvalRequired'
    assert db.get(Client, patients[0].id).notes is None


def test_payment_reports_do_not_repeat_invoice_totals(clinic):
    db, auth, patients, _ = clinic
    from models import BillingPayment
    db.query(BillingPayment).first().currency = 'USD'; db.commit()
    result = reports.preview(db, auth, {'dataset': 'payments', 'aggregates': [{'field': 'amount', 'operation': 'sum'}]})
    assert result['rows'] == [{'currency': 'USD', 'sum_amount': 100}]
    with pytest.raises(HTTPException):
        reports.preview(db, auth, {'dataset': 'payments', 'aggregates': [{'field': 'total', 'operation': 'sum'}]})


def test_report_key_cannot_return_a_different_definition(clinic):
    db, auth, patients, _ = clinic
    reports.queue_export(db, auth, definition(), 'same-request')
    with pytest.raises(HTTPException) as error:
        reports.queue_export(db, auth, definition(columns=['id']), 'same-request')
    assert error.value.detail == 'agent.requestKeyConflict'


def test_clinic_tasks_without_patient_are_reportable(clinic):
    db, auth, patients, _ = clinic
    db.add(FollowupTask(clinic_id=auth.clinic_id, assigned_user_id=auth.user.id, created_by_user_id=auth.user.id,
        title='Clinic stock review', due_date=date(2026, 10, 15)));db.commit()
    result = reports.preview(db, auth, {'dataset': 'tasks', 'columns': ['title', 'client_id']})
    assert result['rows'] == [{'title': 'Clinic stock review', 'client_id': None}]


def test_deleted_record_cannot_be_saved_by_an_approved_draft(clinic):
    db, auth, patients, _ = clinic
    row = db.query(Order).first()
    proposal = changes.prepare(db, auth, 'orders', 'update', {'client_id': patients[0].id, 'type': 'glasses'}, row.id)
    changes.approve(db, auth, proposal['change_id'])
    row.deleted_at = now();db.commit()
    with pytest.raises(HTTPException) as error:
        changes.commit(db, auth, proposal['change_id'])
    assert error.value.status_code == 404


def test_new_exam_measurements_are_visible_and_idempotent(clinic):
    db, auth, patients, _ = clinic
    exam = OpticalExam(client_id=patients[0].id, clinic_id=auth.clinic_id, user_id=auth.user.id, exam_date=date(2026, 10, 14))
    layout = ExamLayout(clinic_id=auth.clinic_id, name='Chosen template', layout_data='{}')
    db.add_all([exam, layout]);db.commit()
    draft = changes.prepare(db, auth, 'exam_measurements', 'create', {'client_id': patients[0].id, 'exam_id': exam.id,
        'layout_id': layout.id, 'exam_data': {'final-prescription': {'r_sph': -1.5}}})
    changes.approve(db, auth, draft['change_id'])
    receipt = changes.commit(db, auth, draft['change_id'])
    assert changes.commit(db, auth, draft['change_id']) == receipt
    assert db.query(ExamLayoutInstance).filter_by(exam_id=exam.id).count() == 1
    assert records.history(db, auth, patients[0].id)['exam_layouts'][0]['data']['final-prescription']['r_sph'] == -1.5


def test_monthly_finance_and_related_patient_age(clinic):
    db, auth, patients, _ = clinic
    result = reports.preview(db, auth, {'dataset': 'billing', 'group_by': ['order_date.month'],
        'aggregates': [{'field': 'outstanding', 'operation': 'sum'}],
        'filters': [{'field': 'client.age', 'operator': 'lt', 'value': 18}], 'as_of': '2026-10-14'})
    assert result['rows'] == [{'order_date.month': 202610, 'currency': 'ILS', 'sum_outstanding': 400}]


def test_order_preview_preserves_unchanged_prescription_and_provenance(clinic):
    db, auth, patients, _ = clinic
    row = db.query(Order).first()
    row.order_data = {'final-prescription': {'r_sph': -1, 'l_sph': -2}, 'legacy_source': {'system': 'fixture'}}
    db.commit()
    draft = changes.prepare(db, auth, 'orders', 'update', {'client_id': patients[0].id,
        'order_data': {'final-prescription': {'r_sph': -1.5}, 'legacy_source': {'system': 'changed'}}}, row.id)
    assert draft['after']['order_data'] == {'final-prescription': {'r_sph': -1.5, 'l_sph': -2},
        'legacy_source': {'system': 'fixture'}}
    changes.approve(db, auth, draft['change_id'])
    changes.commit(db, auth, draft['change_id'])
    db.refresh(row)
    assert row.order_data == draft['after']['order_data']


def test_measurement_preview_uses_same_data_as_commit(clinic):
    db, auth, patients, _ = clinic
    exam = OpticalExam(client_id=patients[0].id, clinic_id=auth.clinic_id, user_id=auth.user.id, exam_date=date(2026, 10, 14))
    layout = ExamLayout(clinic_id=auth.clinic_id, name='Template', layout_data='{}')
    db.add_all([exam, layout]);db.flush()
    instance = ExamLayoutInstance(exam_id=exam.id, layout_id=layout.id, layout_data='{}', exam_data={
        'final-prescription': {'r_sph': -1}, 'subjective': {'r_sph': -2}})
    db.add(instance);db.commit()
    draft = changes.prepare(db, auth, 'exam_measurements', 'update', {'client_id': patients[0].id, 'exam_id': exam.id,
        'layout_instance_id': instance.id, 'exam_data': {'final-prescription': {'r_sph': -1.5}}})
    assert draft['before']['exam_data']['final-prescription']['r_sph'] == -1
    assert draft['after']['exam_data']['subjective']['r_sph'] == -2
    changes.approve(db, auth, draft['change_id']);changes.commit(db, auth, draft['change_id'])
    db.refresh(instance)
    assert instance.exam_data == draft['after']['exam_data']


def test_durable_principal_requires_an_active_registered_clinic_session(clinic):
    db, auth, patients, _ = clinic
    from agent.authority import register_principal
    register_principal(db, auth, 'trusted-customer')
    signed = {'external_user_id': 'trusted-customer', 'chat_scope': f'clinic:{auth.clinic_id}:patient:0'}
    assert resolve_artifact(db, signed).user.id == auth.user.id
    for other in [{'external_user_id': 'foreign', 'chat_scope': signed['chat_scope']},
                  {'external_user_id': 'trusted-customer', 'chat_scope': f'clinic:{auth.clinic_id}:patient:{patients[0].id}'},
                  {'external_user_id': 'trusted-customer', 'chat_scope': 'clinic:999:patient:0'}]:
        with pytest.raises(HTTPException):
            resolve_artifact(db, other)
    db.query(AuthSession).first().revoked_at = now();db.commit()
    with pytest.raises(HTTPException) as error:
        resolve_artifact(db, signed)
    assert error.value.status_code == 401


def test_relogin_retains_owned_proposals_but_rechecks_staff_authority(clinic):
    db, auth, patients, _ = clinic
    from agent.models import AgentContext
    from agent.authority import register_principal
    draft = changes.prepare(db, auth, 'clients', 'update', {'notes': 'Reviewed after login'}, patients[0].id)
    db.query(AuthSession).first().revoked_at = now()
    session = AuthSession(id='00000000-0000-0000-0000-000000000002', user_id=auth.user.id,
        company_id=auth.user.company_id, clinic_id=auth.clinic_id, refresh_token_hash='new-synthetic', expires_at=now()+timedelta(days=1))
    db.add(session);db.flush()
    binding = AgentContext(session_id=session.id, user_id=auth.user.id, company_id=auth.user.company_id,
        clinic_id=auth.clinic_id, patient_scope=0)
    db.add(binding);db.commit()
    fresh = Authority(auth.user, auth.clinic_id, binding.id)
    register_principal(db, fresh, 'trusted-customer')
    signed = {'external_user_id': 'trusted-customer', 'chat_scope': f'clinic:{auth.clinic_id}:patient:0'}
    active = resolve_artifact(db, signed)
    assert active.context_id == fresh.context_id
    changes.approve(db, fresh, draft['change_id'])
    assert changes.commit(db, active, draft['change_id'])['notes'] == 'Reviewed after login'
    hidden = Authority(auth.user, auth.clinic_id, binding.id, patient_id=patients[1].id)
    assert changes.owned_change(db, hidden, draft['change_id']) is None
    session.revoked_at = now();db.commit()
    with pytest.raises(HTTPException):
        resolve_artifact(db, signed)
