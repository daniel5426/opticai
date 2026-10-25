from datetime import date
import json
import pytest
from fastapi import HTTPException
from models import Client, OpticalExam, ExamLayoutInstance, MedicalLog, PrescriptionSearchIndex, BillingPayment, Billing
from services.prescription_search_index import rebuild_exam_instance_index
from services.prescription_matching import search as prescriptions, distance
from services.clinic_query import FilterGroup
from agent import search, report_definitions
from agent.report_runs import materialize, page, paths
from agent.models import AgentJob, SavedReport, PrescriptionIndexState
from agent.config import now


def exam(db, client, data, dated=date(2026, 10, 1)):
    parent = OpticalExam(client_id=client.id, clinic_id=client.clinic_id, exam_date=dated, test_name='Synthetic exam')
    db.add(parent); db.flush()
    instance = ExamLayoutInstance(exam_id=parent.id, exam_data=data)
    db.add(instance); db.flush()
    rebuild_exam_instance_index(db, instance); db.commit()
    if not db.get(PrescriptionIndexState, client.clinic_id):
        db.add(PrescriptionIndexState(clinic_id=client.clinic_id, ready=True, checkpoint={}))
        db.commit()
    return parent, instance


def test_prescriptions_group_rank_equivalence_and_sources(clinic):
    db, auth, patients, _ = clinic
    parent, _ = exam(db, patients[0], {'final-prescription': {'r_sph': -1, 'r_cyl': 1, 'r_ax': 90, 'l_sph': -2}})
    exam(db, patients[1], {'final-prescription': {'r_sph': .25, 'r_cyl': -1, 'r_ax': 0, 'l_sph': -2}})
    exam(db, patients[2], {'final-prescription': {'r_sph': 0, 'r_cyl': -1, 'r_ax': 0}})
    result = prescriptions(db, [auth.clinic_id], {'right': {'sph': 0., 'cyl': -1., 'ax': 180}, 'left': {'sph': -2.}})
    assert result['index_ready']
    assert [r['client_id'] for r in result['items']] == [patients[0].id, patients[1].id]
    first = result['items'][0]['matches'][0]
    assert first['distance'] == 0 and first['eyes']['R']['values']['cyl'] == 1
    assert first['source']['path'].endswith(f'/exams/{parent.id}')


def test_prescriptions_never_pair_cards_or_treat_missing_as_zero(clinic):
    db, auth, patients, _ = clinic
    exam(db, patients[0], {'final-prescription-a': {'r_sph': 0.}, 'final-prescription-b': {'l_sph': -2.}})
    result = prescriptions(db, [auth.clinic_id], {'right': {'sph': 0.}, 'left': {'sph': -2.}})
    assert result['total'] == 0
    assert prescriptions(db, [auth.clinic_id], {'right': {'cyl': 0.}})['total'] == 0
    assert prescriptions(db, [auth.clinic_id], {'right': {'sph': 0.}})['total'] == 1


def test_prescriptions_exclude_deleted_merged_foreign_and_unready(clinic):
    db, auth, patients, _ = clinic
    parent, _ = exam(db, patients[0], {'final-prescription': {'r_sph': -1.}})
    exam(db, patients[1], {'final-prescription': {'r_sph': -1.}})
    assert prescriptions(db, [auth.clinic_id], {'right': {'sph': -1.}}, patients[1].id)['total'] == 1
    parent.deleted_at = now(); patients[1].merged_into_client_id = patients[0].id; db.commit()
    assert prescriptions(db, [auth.clinic_id], {'right': {'sph': -1.}})['total'] == 0


def test_prescription_distance_partial_axis_and_unknown():
    assert distance({'ax': 179}, {'ax': 1})[0] == .2
    assert distance({'sph': 0., 'cyl': 0., 'ax': 40}, {'sph': 0., 'cyl': 0., 'ax': 150})[0] == 0
    assert distance({'sph': 0.}, {'sph': None}) is None
    assert distance({'sph': -1.}, {'sph': -1., 'cyl': 1., 'ax': 90})[0] == 0


def test_search_all_records_and_private_fields(clinic):
    db, auth, patients, _ = clinic
    db.add(MedicalLog(client_id=patients[0].id, clinic_id=auth.clinic_id, log='Synthetic discomfort', log_date=date(2026, 10, 1)))
    patients[0].hidden_note = 'private-needle'; db.commit()
    result = search.search(db, auth, {'text': 'discomfort'})
    assert result['total'] == 1 and result['items'][0]['dataset'] == 'medical_logs'
    assert result['items'][0]['source']['path'].endswith('?tab=medical')
    assert search.search(db, auth, {'text': 'private-needle'})['total'] == 0
    assert search.search(db, auth, {'text': 'Foreign'})['total'] == 0


def test_nested_relationship_report_preserves_patient_grain(clinic):
    db, auth, patients, _ = clinic
    exam(db, patients[0], {'final-prescription': {'r_sph': -1.}})
    exam(db, patients[0], {'final-prescription': {'r_sph': -2.}})
    raw = {'version': 2, 'dataset': 'clients', 'columns': ['id', 'first_name'], 'where': {'match': 'all', 'conditions': [
        {'field': 'exams.exam_date', 'operator': 'gte', 'value': '2026-01-01'},
        {'match': 'any', 'conditions': [{'field': 'first_name', 'operator': 'eq', 'value': 'Boy'},
            {'field': 'first_name', 'operator': 'eq', 'value': 'Nobody'}]}]}}
    definition = report_definitions.resolve(raw)
    query, columns = report_definitions.compile_definition(db, auth, definition)
    assert query.count() == 1 and columns == ['id', 'first_name']
    with pytest.raises(HTTPException):
        report_definitions.compile_definition(db, auth, report_definitions.resolve({**raw, 'where': {'conditions': [{'field': 'hidden_note', 'operator': 'eq', 'value': 'secret'}]}}))


def test_snapshot_paging_exports_survive_data_changes(clinic):
    db, auth, patients, _ = clinic
    for i in range(30):
        db.add(Client(clinic_id=auth.clinic_id, company_id=auth.user.company_id, first_name=f'Snapshot-{i}'))
    db.commit()
    queued = report_definitions.queue(db, auth, {'version': 2, 'dataset': 'clients', 'columns': ['id', 'first_name'],
        'title': 'Synthetic report'}, 'snapshot-test')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    assert job.status == 'completed' and job.result['total'] == 32
    assert len(page(job)['rows']) == 25 and len(page(job, 25)['rows']) == 7
    original = page(job)['rows']
    patients[0].first_name = 'Changed after run'; db.commit()
    materialize(db, auth, job)
    assert page(job)['rows'] == original
    assert all(path.exists() for path in paths(job))
    from openpyxl import load_workbook
    workbook = load_workbook(paths(job)[1], read_only=True)
    assert sum(1 for _ in workbook['Records'].rows) == 33
    workbook.close()
    assert report_definitions.queue(db, auth, job.definition, 'snapshot-test')['job_id'] == job.id


def test_saved_reports_private_scoped_revision_and_dynamic_defaults(clinic):
    db, auth, patients, _ = clinic
    saved = report_definitions.saved(db, auth, 'save', title='Health', definition={'version': 2, 'preset': 'business_health'})
    assert 'end_date' not in saved['definition']
    assert report_definitions.saved(db, auth, 'list')['items'][0]['id'] == saved['id']
    with pytest.raises(HTTPException):
        report_definitions.saved(db, auth, 'rename', saved['id'], title='New', revision=9)
    changed = report_definitions.saved(db, auth, 'rename', saved['id'], title='New', revision=1)
    assert changed['revision'] == 2
    auth.patient_id = patients[0].id
    assert report_definitions.saved(db, auth, 'list')['items'] == []
    with pytest.raises(HTTPException):
        report_definitions.saved(db, auth, 'delete', saved['id'], revision=2)


def test_business_health_refunds_no_shared_invoice_duplication(clinic):
    db, auth, patients, _ = clinic
    paid = db.query(BillingPayment).one()
    db.add(BillingPayment(billing_id=paid.billing_id, amount=20, kind='refund', currency='ILS', paid_at=date(2026, 10, 3)))
    db.get(Billing, paid.billing_id).prepayment_amount = 80
    db.commit()
    result = report_definitions.queue(db, auth, {'version': 2, 'preset': 'business_health',
        'start_date': '2026-10-01', 'end_date': '2026-10-31'}, 'health')
    job = db.get(AgentJob, result['job_id']); materialize(db, auth, job)
    metrics = {item['key']: item['current'] for item in job.result['metrics'] if item['currency'] == 'ILS'}
    assert metrics['sales'] == 500 and metrics['collected'] == 80 and metrics['billed_orders'] == 1
    assert metrics['outstanding'] == 420
    assert sum(row['sales'] for row in page(job)['rows']) == 500


def test_index_backfill_resumes_and_preserves_other_sources(clinic):
    from services.prescription_index_maintenance import backfill
    db, auth, patients, _ = clinic
    parent, instance = exam(db, patients[0], {'final-prescription-one': {'r_sph': -1., 'l_sph': -2.}})
    state = db.get(PrescriptionIndexState, auth.clinic_id)
    state.ready = False
    db.query(PrescriptionSearchIndex).filter_by(clinic_id=auth.clinic_id).update({'component_key': None})
    db.commit()
    assert not prescriptions(db, [auth.clinic_id], {'right': {'sph': -1.}})['index_ready']
    first = backfill(db, auth.clinic_id, batch_size=1, max_batches=1)
    assert not first['ready']
    assert db.query(PrescriptionSearchIndex).filter_by(source_type='exam', source_id=instance.id).count() == 2
    assert backfill(db, auth.clinic_id, batch_size=1)['ready']
    assert prescriptions(db, [auth.clinic_id], {'right': {'sph': -1.}})['total'] == 1


def test_same_record_relationship_group_and_measurement_reporting(clinic):
    db, auth, patients, _ = clinic
    parent, _ = exam(db, patients[0], {'final-prescription': {'r_sph': -1., 'l_sph': -2.}}, date(2026, 10, 1))
    parent.test_name = 'First'; db.commit()
    other, _ = exam(db, patients[0], {'final-prescription': {'r_sph': -3.}}, date(2026, 9, 1))
    other.test_name = 'Second'; db.commit()
    definition = report_definitions.resolve({'version': 2, 'dataset': 'clients', 'columns': ['id'], 'where': {
        'relation': 'exams', 'conditions': [{'field': 'exam_date', 'operator': 'eq', 'value': '2026-10-01'},
            {'field': 'test_name', 'operator': 'eq', 'value': 'Second'}]}})
    assert report_definitions.compile_definition(db, auth, definition)[0].count() == 0
    definition = report_definitions.resolve({'version': 2, 'dataset': 'measurements', 'group_by': ['eye'],
        'aggregates': [{'field': 'client_id', 'operation': 'count_distinct'}]})
    query, columns = report_definitions.compile_definition(db, auth, definition)
    assert len(query.all()) == 2 and 'count_distinct_client_id' in columns


def test_period_comparison_and_saved_relative_business_report(clinic):
    db, auth, patients, _ = clinic
    queued = report_definitions.queue(db, auth, {'version': 2, 'dataset': 'payments', 'date_field': 'paid_at',
        'start_date': '2026-10-01', 'end_date': '2026-10-31', 'aggregates': [{'field': 'amount', 'operation': 'sum'}]}, 'comparison')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    assert job.result['comparison']['start'] == '2026-08-31'
    from openpyxl import load_workbook
    book = load_workbook(paths(job)[1], read_only=True)
    assert 'Previous period' in book.sheetnames; book.close()
    definition = report_definitions.resolve({'version': 2, 'preset': 'business_health'}).model_dump(mode='json')
    saved = report_definitions.saved(db, auth, 'save', title='Rolling health', definition=definition)
    assert saved['definition']['relative_period'] and 'end_date' not in saved['definition']


def test_legacy_prescription_contract_prebackfill_never_combines_cards(clinic):
    from schemas import PrescriptionSearchRequest
    from services.prescription_matching import legacy_search
    db, auth, patients, _ = clinic
    exam(db, patients[0], {'final-prescription-a': {'r_sph': -1.}, 'final-prescription-b': {'l_sph': -2.}})
    db.query(PrescriptionSearchIndex).update({'component_key': None}); db.commit()
    request = PrescriptionSearchRequest(right={'sph': -1.}, left={'sph': -2.})
    assert legacy_search(db, [auth.clinic_id], request)['total'] == 0
    request = PrescriptionSearchRequest(right={'sph': -1.})
    result = legacy_search(db, [auth.clinic_id], request)
    assert result['total'] == 1 and result['items'][0]['matched_eyes'] == ['R']


def test_equal_weight_fields_across_eyes_and_measurement_readiness(clinic):
    db, auth, patients, _ = clinic
    exam(db, patients[0], {'final-prescription': {'r_sph': -1., 'l_sph': -2., 'l_add': 1.}})
    result = prescriptions(db, [auth.clinic_id], {'right': {'sph': -1.25}, 'left': {'sph': -2., 'add': 1.}})
    assert result['items'][0]['matches'][0]['distance'] == pytest.approx(1 / 3, abs=1e-6)
    db.get(PrescriptionIndexState, auth.clinic_id).ready = False; db.commit()
    definition = report_definitions.resolve({'dataset': 'clients', 'columns': ['id'], 'where': {
        'conditions': [{'field': 'measurements.sph', 'operator': 'eq', 'value': -1.}]}})
    with pytest.raises(HTTPException) as error:
        report_definitions.compile_definition(db, auth, definition)
    assert error.value.status_code == 409
    assert distance({'sph': -1., 'cyl': 0., 'ax': 0}, {'sph': -1., 'cyl': 0., 'ax': None})[0] == 0


def test_report_http_contract_exports_expiry_and_scope(clinic):
    from datetime import timedelta
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from EndPoints import agent as endpoints
    db, auth, patients, _ = clinic
    queued = report_definitions.queue(db, auth, {'dataset': 'clients', 'columns': ['id', 'first_name']}, 'http-report')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    app = FastAPI(); app.include_router(endpoints.router, prefix='/api/v1')
    app.dependency_overrides[endpoints.authorized] = lambda: auth
    app.dependency_overrides[endpoints.get_db] = lambda: db
    prefix = f'/api/v1/agent/reports/{job.id}'
    with TestClient(app) as client:
        metadata = client.get(f'/api/v1/agent/jobs/{job.id}?metadata_only=true').json()
        assert 'result' not in metadata and metadata['download_path'].endswith('/download')
        assert metadata['title'] == job.definition['title']
        assert client.get(prefix + '/page?limit=1').json()['next_offset'] == 1
        assert client.get(prefix + '/page?offset=-1').status_code == 422
        assert client.get(prefix + '/export?format=csv').status_code == 200
        assert client.get(prefix + '/export?format=sql').status_code == 422
        saved = client.post('/api/v1/agent/saved-reports', json={'action': 'save', 'title': 'Private',
            'definition': {'version': 2, 'dataset': 'clients', 'columns': ['id']}})
        assert saved.status_code == 200
        assert len(client.get('/api/v1/agent/saved-reports').json()['items']) == 1
        auth.patient_id = patients[0].id
        assert client.get(prefix + '/page').status_code == 403
        assert client.get(f'/api/v1/agent/jobs/{job.id}?metadata_only=true').status_code == 403
        assert client.get('/api/v1/agent/saved-reports').json()['items'] == []
        auth.patient_id = None
        job.expires_at = now() - timedelta(seconds=1); db.commit()
        assert client.get(prefix + '/page').status_code == 410
        assert client.get(prefix + '/export').status_code == 410
        assert client.get('/api/v1/agent/jobs/' + job.id).status_code == 410


def test_search_pagination_uses_stable_relevance_order(clinic):
    db, auth, patients, _ = clinic
    for i in range(35):
        db.add(Client(clinic_id=auth.clinic_id, company_id=auth.user.company_id,
            first_name='Remember', last_name='Remember' if i > 28 else f'Patient {i}'))
    db.commit()
    pages = [search.search(db, auth, {'text': 'Remember', 'datasets': ['clients'], 'offset': offset, 'limit': 10})
        for offset in (0, 10, 20, 30)]
    ids = [item['values']['id'] for result in pages for item in result['items']]
    assert len(ids) == len(set(ids)) == 35
    assert all(len(item['matched_fields']) == 2 for item in pages[0]['items'][:6])


def test_timestamp_date_filter_includes_full_israel_local_day(clinic):
    from datetime import datetime
    from models import File
    db, auth, patients, _ = clinic
    for hour in (18, 21):
        db.add(File(clinic_id=auth.clinic_id, client_id=patients[0].id, file_name=str(hour),
            upload_date=datetime(2026, 10, 1, hour), file_type='text/plain', file_size=10))
    db.commit()
    definition = report_definitions.resolve({'dataset': 'files', 'columns': ['file_name'],
        'date_field': 'upload_date', 'start_date': '2026-10-01', 'end_date': '2026-10-01'})
    query, _ = report_definitions.compile_definition(db, auth, definition)
    assert [row[0] for row in query] == ['18']  # 21:00 UTC is already the next Israel-local day.


def test_derived_metrics_zero_denominators_and_currency_charts(clinic):
    db, auth, patients, _ = clinic
    original = db.query(Billing).one()
    from models import Order
    order = Order(clinic_id=auth.clinic_id, client_id=patients[0].id, order_date=date(2026, 10, 2))
    db.add(order); db.flush()
    db.add(Billing(order_id=order.id, total_after_discount=200, prepayment_amount=0, currency='USD'))
    db.commit()
    definition = {'dataset': 'billing', 'group_by': ['type'], 'aggregates': [
        {'operation': 'sum', 'field': 'total'}, {'operation': 'count_distinct', 'field': 'id'}, {'operation': 'sum', 'field': 'paid'}],
        'derived_metrics': [{'key': 'aov', 'label': 'Average order value', 'operation': 'ratio', 'numerator': 'sum_total', 'denominator': 'count_distinct_id'},
            {'key': 'ratio', 'label': 'Zero-safe ratio', 'operation': 'ratio', 'numerator': 'sum_total', 'denominator': 'sum_paid'}]}
    queued = report_definitions.queue(db, auth, definition, 'derived')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    rows = {row['currency']: row for row in page(job)['rows']}
    assert rows['ILS']['aov'] == 500 and rows['USD']['aov'] == 200 and rows['USD']['ratio'] is None
    assert {chart['currency'] for chart in job.result['charts']} == {'ILS', 'USD'}
    assert all(len({row['currency'] for row in chart['rows']}) == 1 for chart in job.result['charts'])
    assert job.result['column_labels']['aov'] == 'Average order value'
    with pytest.raises(HTTPException):
        report_definitions.compile_definition(db, auth, report_definitions.resolve({**definition,
            'derived_metrics': [{'key': 'bad', 'label': 'Bad', 'operation': 'ratio', 'numerator': 'unsafe', 'denominator': 'sum_total'}]}))


def test_presentation_uses_snapshot_columns_and_keeps_currencies_separate(clinic):
    db, auth, patients, _ = clinic
    definition = {'dataset': 'payments', 'group_by': ['kind'], 'aggregates': [{'operation': 'sum', 'field': 'amount'}],
        'title': 'Collections by method', 'presentation': {'show_table': False, 'labels': {'sum_amount': 'Collections'},
            'charts': [{'kind': 'donut', 'title': 'Payment mix', 'x': 'kind', 'series': ['sum_amount']}]}}
    queued = report_definitions.queue(db, auth, definition, 'visuals')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    assert not job.result['show_table'] and job.result['column_labels']['sum_amount'] == 'Collections'
    assert job.result['charts'][0]['kind'] == 'donut' and job.result['charts'][0]['currency'] == 'ILS'
    assert job.result['charts'][0]['rows'][0]['sum_amount'] == 100
    assert 'metrics' not in job.result  # Grouped distinct/count/average values are never added into fabricated KPIs.
    before = job.result['charts']
    db.query(BillingPayment).one().amount = 999; db.commit()
    materialize(db, auth, job)
    assert job.result['charts'] == before
    for invalid in ({'charts': [{'kind': 'bar', 'x': 'kind', 'series': ['private_field']}]}, {'charts': [{'kind': 'bar', 'x': 'kind', 'series': ['currency']}]}, {'labels': {'private_field': 'Bad'}}):
        with pytest.raises(HTTPException) as error:
            report_definitions.queue(db, auth, {**definition, 'presentation': invalid}, 'invalid')
        assert error.value.status_code == 422


def test_single_aggregate_rows_generate_honest_stat_cards_and_old_retries_work(clinic):
    db, auth, *_ = clinic
    definition = {'dataset': 'billing', 'aggregates': [{'operation': 'sum', 'field': 'total'}, {'operation': 'sum', 'field': 'outstanding'}],
        'presentation': {'labels': {'sum_total': 'Sales', 'sum_outstanding': 'Outstanding'}}}
    queued = report_definitions.queue(db, auth, definition, 'stat-cards')
    job = db.get(AgentJob, queued['job_id']); materialize(db, auth, job)
    metrics = {metric['key']: metric['current'] for metric in job.result['metrics']}
    assert metrics == {'sum_total': 500, 'sum_outstanding': 400}
    old = report_definitions.queue(db, auth, {'dataset': 'clients', 'columns': ['id']}, 'old-retry')
    old_job = db.get(AgentJob, old['job_id'])
    old_job.definition = {key: value for key, value in old_job.definition.items() if key != 'presentation'}; db.commit()
    assert report_definitions.queue(db, auth, {'dataset': 'clients', 'columns': ['id']}, 'old-retry')['job_id'] == old_job.id
