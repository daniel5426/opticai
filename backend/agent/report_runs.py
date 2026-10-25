"""Materialize once, then page and export the exact same report execution."""
import csv
import json
import sqlite3
from pathlib import Path
from fastapi import HTTPException
from .config import data_directory, now
from .records import json_value
from .report_definitions import resolve, compile_definition


def paths(job):
    root = data_directory()
    return root / (job.id + '.snapshot.sqlite'), root / (job.id + '.xlsx'), root / (job.id + '.csv')


def temporary_paths(job):
    snapshot, excel, csv_path = paths(job)
    return snapshot.with_suffix('.part'), excel.with_suffix('.xlsx.part'), csv_path.with_suffix('.csv.part')


def materialize(db, authority, job):
    from .reports import safe_csv
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from .report_definitions import ReportDefinitionV2
    definition = ReportDefinitionV2.model_validate(job.definition)
    repeatable = db.get_bind().dialect.name == 'postgresql'
    def cancelled():
        if repeatable:
            # A separate connection observes cancellation while source reads retain
            # their repeatable-read snapshot.
            from sqlalchemy.orm import Session
            from .models import AgentJob
            with Session(bind=db.get_bind()) as current:
                return current.query(AgentJob.status).filter_by(id=job.id).scalar() == 'cancelled'
        db.refresh(job)
        return job.status == 'cancelled'
    snapshot, excel, csv_path = paths(job)
    metadata = {}
    if not snapshot.exists():
        if definition.preset == 'business_health':
            from .business_health import generate
            metadata = generate(db, authority, definition)
            columns, rows = metadata['columns'], metadata.pop('rows')
        else:
            query, columns = compile_definition(db, authority, definition)
            rows = ({key: json_value(value) for key, value in zip(columns, row)} for row in query.yield_per(500))
        temp = snapshot.with_suffix('.part')
        with sqlite3.connect(temp) as handle:
            handle.execute('CREATE TABLE IF NOT EXISTS rows (position INTEGER PRIMARY KEY, payload TEXT NOT NULL)')
            handle.execute('DELETE FROM rows')
            count = 0
            for position, row in enumerate(rows):
                if position % 500 == 0:
                    if cancelled():
                        break
                handle.execute('INSERT INTO rows VALUES (?, ?)', (position, json.dumps(row, ensure_ascii=False)))
                count += 1
            if cancelled():
                temp.unlink(missing_ok=True)
                return
            metadata.update(columns=columns, total=count, generated_at=now().isoformat(), definition=job.definition,
                column_labels={metric.key: metric.label for metric in definition.derived_metrics})
            if definition.preset == 'custom' and definition.date_field and definition.comparison:
                from datetime import timedelta
                days = (definition.end_date - definition.start_date).days + 1
                previous = definition.model_copy(update={'start_date': definition.start_date - timedelta(days=days),
                    'end_date': definition.start_date - timedelta(days=1)})
                previous_query, _ = compile_definition(db, authority, previous)
                handle.execute('CREATE TABLE IF NOT EXISTS comparison_rows (position INTEGER PRIMARY KEY, payload TEXT NOT NULL)')
                handle.execute('DELETE FROM comparison_rows')
                previous_count = 0
                for position, row in enumerate(previous_query.yield_per(500)):
                    if position % 500 == 0 and cancelled():
                        temp.unlink(missing_ok=True)
                        return
                    handle.execute('INSERT INTO comparison_rows VALUES (?, ?)', (position, json.dumps(
                        {key: json_value(value) for key, value in zip(columns, row)}, ensure_ascii=False)))
                    previous_count += 1
                metadata['comparison'] = {'start': previous.start_date.isoformat(), 'end': previous.end_date.isoformat(),
                    'total': previous_count, 'rows': [json.loads(row[0]) for row in handle.execute('SELECT payload FROM comparison_rows ORDER BY position LIMIT 25')]}
            from .report_presentation import apply_presentation, validate_presentation
            validate_presentation(definition, columns)
            apply_presentation(handle, definition, columns, metadata)
            handle.execute('CREATE TABLE IF NOT EXISTS metadata (payload TEXT NOT NULL)')
            handle.execute('DELETE FROM metadata')
            handle.execute('INSERT INTO metadata VALUES (?)', (json.dumps(metadata, ensure_ascii=False),))
        temp.chmod(0o600); temp.replace(snapshot)
    with sqlite3.connect(f'file:{snapshot}?mode=ro', uri=True) as handle:
        metadata = json.loads(handle.execute('SELECT payload FROM metadata').fetchone()[0])
        columns = metadata['columns']
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('Records'); sheet.append(columns)
        csv_temp, excel_temp = csv_path.with_suffix('.csv.part'), excel.with_suffix('.xlsx.part')
        with csv_temp.open('w', encoding='utf-8-sig', newline='') as output:
            extra = ['metric', 'current', 'previous', 'change_percent'] if metadata.get('metrics') else []
            csv_columns = list(dict.fromkeys(columns + extra))
            sections = bool(metadata.get('comparison') or metadata.get('metrics'))
            writer = csv.writer(output); writer.writerow((['section'] if sections else []) + csv_columns)
            for position, record in handle.execute('SELECT position, payload FROM rows ORDER BY position'):
                if position % 500 == 0:
                    if cancelled():
                        workbook.close(); csv_temp.unlink(missing_ok=True); return
                row = json.loads(record)
                values = [row.get(key) for key in columns]
                writer.writerow((['records'] if sections else []) + [safe_csv(row.get(key)) for key in csv_columns])
                if position and position % 1_048_575 == 0:
                    sheet = workbook.create_sheet(f'Records {position // 1_048_575 + 1}'); sheet.append(columns)
                cells = []
                for value in values:
                    cell = WriteOnlyCell(sheet, value=json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cells.append(cell)
                sheet.append(cells)
            if metadata.get('comparison'):
                previous_sheet = workbook.create_sheet('Previous period'); previous_sheet.append(columns)
                for record, in handle.execute('SELECT payload FROM comparison_rows ORDER BY position'):
                    row = json.loads(record)
                    writer.writerow(['previous'] + [safe_csv(row.get(key)) for key in csv_columns])
                    previous_sheet.append([safe_csv(row.get(key)) for key in columns])
            for metric in metadata.get('metrics', []):
                writer.writerow(['metrics'] + [safe_csv(metric['key'] if key == 'metric' else metric.get(key)) for key in csv_columns])
        definitions = workbook.create_sheet('Definition')
        definitions.append(['Definition', json.dumps(job.definition, ensure_ascii=False)])
        definitions.append(['Generated UTC', metadata['generated_at']])
        for key, value in metadata.get('definitions', {}).items():
            definitions.append([key, value])
        if metadata.get('metrics'):
            metrics = workbook.create_sheet('Metrics'); metrics.append(['Metric', 'Currency', 'Current', 'Previous', 'Change %'])
            for item in metadata['metrics']:
                metrics.append([item['key'], item['currency'], item['current'], item['previous'], item['change_percent']])
        workbook.save(excel_temp)
        for original, target in ((csv_temp, csv_path), (excel_temp, excel)):
            original.chmod(0o600); original.replace(target)
    if repeatable:
        db.rollback()
    db.refresh(job)
    if job.status == 'cancelled':
        return
    job.file_name = excel.name if definition.format == 'xlsx' else csv_path.name
    job.result = {**metadata, 'type': 'report_result', 'job_id': job.id, **page(job, 0, 25),
        'exports': {format: f'/api/v1/agent/reports/{job.id}/export?format={format}' for format in ('xlsx', 'csv')}}
    job.status, job.updated_at = 'completed', now()
    db.commit()


def page(job, offset=0, limit=25):
    path = paths(job)[0]
    if not path.exists():
        raise HTTPException(410, 'agent.downloadUnavailable')
    with sqlite3.connect(f'file:{path}?mode=ro', uri=True) as handle:
        metadata = json.loads(handle.execute('SELECT payload FROM metadata').fetchone()[0])
        rows = [json.loads(row[0]) for row in handle.execute('SELECT payload FROM rows ORDER BY position LIMIT ? OFFSET ?', (limit, offset))]
    return {'rows': rows, 'total': metadata['total'], 'columns': metadata['columns'], 'offset': offset,
        'next_offset': offset + limit if offset + limit < metadata['total'] else None}
