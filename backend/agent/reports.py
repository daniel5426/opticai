"""Validated report definitions compile to scoped ORM queries, never model SQL."""
import csv
import json
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from decimal import Decimal
from typing import Any, Literal
from pydantic import BaseModel, ConfigDict, Field
from fastapi import HTTPException
from sqlalchemy.exc import IntegrityError
from sqlalchemy import select, union_all, func, case, and_, or_, literal
from models import (Client, Billing, Order, ContactLensOrder, OrderGroup, BillingPayment,
                    CatalogProduct, CatalogVariant, InventoryBalance, OpticalExam)
from .records import MODELS, FIELDS, json_value, scoped_query
from .models import AgentJob
from .config import now, data_directory


class Filter(BaseModel):
    model_config = ConfigDict(extra='forbid')
    field: str
    operator: Literal['eq', 'ne', 'lt', 'lte', 'gt', 'gte', 'contains', 'in', 'empty', 'not_empty']
    value: Any = None


class Aggregate(BaseModel):
    model_config = ConfigDict(extra='forbid')
    field: str
    operation: Literal['sum', 'count', 'count_distinct', 'min', 'max', 'avg']


class ReportDefinition(BaseModel):
    model_config = ConfigDict(extra='forbid')
    dataset: Literal['clients', 'appointments', 'exams', 'orders', 'contact_lens_orders',
                     'referrals', 'medical_logs', 'files', 'tasks', 'billing', 'payments', 'inventory']
    columns: list[str] = Field(default_factory=list, max_length=40)
    filters: list[Filter] = Field(default_factory=list, max_length=50)
    match: Literal['all', 'any'] = 'all'
    group_by: list[str] = Field(default_factory=list, max_length=8)
    aggregates: list[Aggregate] = Field(default_factory=list, max_length=12)
    sort_by: list[str] = Field(default_factory=list, max_length=8)
    format: Literal['xlsx', 'csv'] = 'xlsx'
    title: str = Field(default='Clinic report', max_length=200)
    as_of: date = Field(default_factory=lambda: now().astimezone(ZoneInfo('Asia/Jerusalem')).date())


from services.clinic_finance import invoice_ledger as billing_owners
from services.clinic_catalog import expressions, age_at, CLIENT_REPORT_FIELDS


from services.clinic_query import predicate


def compile_report(db, authority, definition):
    if definition.dataset == 'measurements':
        from .models import PrescriptionIndexState
        state = db.get(PrescriptionIndexState, authority.clinic_id)
        if not state or not state.ready:
            raise HTTPException(409, 'agent.prescriptionIndexPending')
    used = list(definition.group_by) + [f.field for f in definition.filters] + [a.field for a in definition.aggregates]
    if not definition.aggregates:
        used += definition.columns or ['total', 'paid', 'outstanding', 'currency']
    def collect(group):
        if group is None:
            return
        if hasattr(group, 'conditions'):
            if not group.relation:
                for child in group.conditions:
                    collect(child)
        else:
            used.append(group.field)
    collect(getattr(definition, 'where', None))
    query, fields = expressions(db, authority, definition.dataset, definition.as_of,
        client_balances=bool(set(used) & {'total', 'paid', 'outstanding', 'currency'}))
    age_key = 'age' if definition.dataset == 'clients' else 'client.age'
    if age_key in fields:
        fields[age_key] = age_at(definition.as_of)
    requested = definition.columns or (['id', 'first_name', 'last_name', 'age', 'gender', 'total', 'paid', 'outstanding', 'currency']
                                      if definition.dataset == 'clients' else sorted(key for key in fields if not key.endswith(('.year', '.month')))[:12])
    for key in [*requested, *definition.group_by, *(f.field for f in definition.filters), *(a.field for a in definition.aggregates)]:
        if key not in fields:
            raise HTTPException(422, 'agent.unsupportedReportField')
    clauses = [predicate(fields[item.field], item) for item in definition.filters]
    if clauses:
        query = query.filter((and_ if definition.match == 'all' else or_)(*clauses))
    selected = {key: fields[key] for key in (definition.group_by if definition.aggregates else (definition.group_by or requested))}
    if definition.aggregates:
        # Amount aggregations always retain currency to avoid fabricated FX totals.
        if any(item.field in {'total', 'paid', 'outstanding', 'amount'} for item in definition.aggregates) and 'currency' in fields:
            selected.setdefault('currency', fields['currency'])
        for item in definition.aggregates:
            if item.operation in {'sum', 'avg'}:
                try:
                    numeric = fields[item.field].type.python_type in {int, float, Decimal}
                except (NotImplementedError, AttributeError):
                    numeric = item.field in {'age', 'client.age', 'amount', 'total', 'paid', 'outstanding', 'available'}
                if not numeric:
                    raise HTTPException(422, 'agent.invalidFields')
            selected[f'{item.operation}_{item.field}'] = (func.count(func.distinct(fields[item.field])) if item.operation == 'count_distinct' or definition.dataset == 'clients' and item.operation == 'count' and item.field == 'id' else getattr(func, item.operation)(fields[item.field]))
    aggregate_keys = {f'{a.operation}_{a.field}' for a in definition.aggregates}
    derived_keys = set()
    for metric in getattr(definition, 'derived_metrics', []):
        if metric.key in selected or metric.key in fields or metric.numerator not in aggregate_keys or metric.denominator not in aggregate_keys:
            raise HTTPException(422, 'agent.invalidFields')
        for aggregate in definition.aggregates:
            if f'{aggregate.operation}_{aggregate.field}' in {metric.numerator, metric.denominator} and aggregate.operation not in {'count', 'count_distinct'}:
                try:
                    numeric = fields[aggregate.field].type.python_type in {int, float, Decimal}
                except (NotImplementedError, AttributeError):
                    numeric = aggregate.field in {'age', 'client.age', 'amount', 'total', 'paid', 'outstanding', 'available'}
                if not numeric:
                    raise HTTPException(422, 'agent.invalidFields')
        numerator, denominator = selected[metric.numerator], selected[metric.denominator]
        if metric.operation == 'difference':
            selected[metric.key] = numerator - denominator
        else:
            selected[metric.key] = numerator * (100.0 if metric.operation == 'percent' else 1.0) / func.nullif(denominator, 0)
        derived_keys.add(metric.key)
    query = query.with_entities(*[expr.label(key) for key, expr in selected.items()])
    if definition.group_by or definition.aggregates:
        query = query.group_by(*[expr for key, expr in selected.items() if key not in aggregate_keys | derived_keys])
    for key in definition.sort_by:
        descending = key.startswith('-')
        key = key.lstrip('-')
        if key not in selected:
            raise HTTPException(422, 'agent.unsupportedReportField')
        query = query.order_by(selected[key].desc() if descending else selected[key].asc())
    if not definition.sort_by:
        query = query.order_by(*[value.asc() for value in selected.values()])
    return query, list(selected)


def catalog(db, authority):
    return {'datasets': {name: sorted(expressions(db, authority, name)[1]) for name in
            [*MODELS, 'billing', 'payments', 'inventory']},
            'definition_schema': ReportDefinition.model_json_schema(),
            'operators': list(Filter.model_fields['operator'].annotation.__args__),
            'notes': 'Date .month fields are numeric YYYYMM periods; .year fields are calendar years. Client referring_party is the recorded referring-agent/party text; clarify ambiguous relationships. Client balances are grouped by currency. Age is as of the report date. Unknown birthdays do not match age filters.'}


def preview(db, authority, raw):
    definition = ReportDefinition.model_validate(raw)
    query, columns = compile_report(db, authority, definition)
    return {'definition': definition.model_dump(mode='json'), 'columns': columns, 'total': query.count(),
            'rows': [{key: json_value(value) for key, value in zip(columns, row)} for row in query.limit(25)],
            'preview_limit': 25}


def queue_export(db, authority, raw, request_key):
    definition = ReportDefinition.model_validate(raw)
    compile_report(db, authority, definition)
    key = f'report:{authority.context_id}:{request_key}'
    job = db.query(AgentJob).filter_by(unique_key=key).first()
    if not job:
        job = AgentJob(unique_key=key, company_id=authority.user.company_id, clinic_id=authority.clinic_id,
            user_id=authority.user.id, context_id=authority.context_id, kind='report',
            definition=definition.model_dump(mode='json'), expires_at=now() + timedelta(days=7))
        db.add(job)
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            job = db.query(AgentJob).filter_by(unique_key=key).one()
    if job.definition != definition.model_dump(mode='json'):
        raise HTTPException(409, 'agent.requestKeyConflict')
    return {'job_id': job.id, 'status': job.status, 'path': f'/api/v1/agent/jobs/{job.id}',
            'download_path': f'/api/v1/agent/jobs/{job.id}/download'}


def export_report(db, authority, job):
    if job.definition.get('version') == 2:
        from .report_runs import materialize
        return materialize(db, authority, job)
    definition = ReportDefinition.model_validate(job.definition)
    query, columns = compile_report(db, authority, definition)
    root = data_directory()
    temporary = root / (job.id + '.part')
    target = root / (job.id + '.' + definition.format)
    count = 0
    if definition.format == 'csv':
        with temporary.open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            for row in query.yield_per(500):
                if count % 500 == 0:
                    db.refresh(job)
                    if job.status == 'cancelled':
                        temporary.unlink(missing_ok=True)
                        return
                writer.writerow([safe_csv(json_value(value)) for value in row])
                count += 1
    else:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('Records')
        sheet.append(columns)
        for row in query.yield_per(500):
            if count % 500 == 0:
                db.refresh(job)
                if job.status == 'cancelled':
                    workbook.close()
                    temporary.unlink(missing_ok=True)
                    return
            if count and count % 1_048_575 == 0:
                sheet = workbook.create_sheet(f'Records {count // 1_048_575 + 1}')
                sheet.append(columns)
            cells = []
            for value in row:
                value = json_value(value)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'
                cells.append(cell)
            sheet.append(cells)
            count += 1
        summary = workbook.create_sheet('Definition')
        summary.append(['Report', definition.title])
        summary.append(['Generated UTC', now().isoformat()])
        summary.append(['Definition', definition.model_dump_json()])
        summary.append(['Records', count])
        workbook.save(temporary)
    job = db.query(AgentJob).filter_by(id=job.id).populate_existing().with_for_update().one()
    if job.status == 'cancelled':
        temporary.unlink(missing_ok=True)
        return
    temporary.replace(target)
    target.chmod(0o600)
    job.file_name = target.name
    job.result = {'rows': count, 'columns': columns, 'generated_at': now().isoformat()}
    job.status = 'completed'
    job.updated_at = now()
    db.commit()


def safe_csv(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value
